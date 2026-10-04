"""Авторизация и полный путь загрузки на настоящем SQL без внешнего сервера."""

import io
from contextlib import asynccontextmanager
from datetime import timedelta
from uuid import uuid4

import docx
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import MetaData, create_engine, event, select
from sqlalchemy.orm import Session

from app.auth import password_hash
from app.clock import get_clock
from app.db import get_session
from app.main import create_app
from app.models import Appointment, Base, Finding, Patient, Protocol, Study
from app.services.anonymization import anonymize
from app.services.mock_mis import MockMisService
from app.settings import get_settings
from tests.conftest import auth_headers

pytestmark = pytest.mark.anyio


class SqlSession:
    """Синхронная SQLite с асинхронным интерфейсом; выполняет запросы боевого кода."""

    def __init__(self):
        self.engine = create_engine("sqlite://")

        @event.listens_for(self.engine, "connect")
        def connect(connection, _):
            connection.create_function("pg_advisory_xact_lock", 1, lambda _: 1)
            connection.execute("PRAGMA foreign_keys=ON")

        metadata = MetaData()
        for table in Base.metadata.sorted_tables:
            copied = table.to_metadata(metadata)
            for column in copied.columns:
                if column.server_default is not None and (
                    "::" in str(column.server_default.arg) or column.name == "id"
                ):
                    column.server_default = None
        metadata.create_all(self.engine)
        self.sync = Session(self.engine, expire_on_commit=False)

    def add(self, row):
        self.sync.add(row)

    @asynccontextmanager
    async def begin_nested(self):
        with self.sync.begin_nested():
            yield

    async def flush(self):
        self.sync.flush()

    async def execute(self, query):
        return self.sync.execute(query)

    async def scalar(self, query):
        return self.sync.scalar(query)

    async def scalars(self, query):
        return self.sync.scalars(query)

    async def get(self, model, identifier):
        return self.sync.get(model, identifier)

    async def commit(self):
        self.sync.commit()

    async def rollback(self):
        self.sync.rollback()


@pytest.fixture
async def secured(model_clock):
    session = SqlSession()
    app = create_app()
    app.dependency_overrides[get_session] = lambda: session
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield client, session
    session.sync.close()
    session.engine.dispose()


def make_patient(session, identifier=None, card="card-123"):
    row = Patient(
        id=identifier or uuid4(),
        external_id=card,
        anonymized_hash=uuid4().hex,
        created_at=get_clock().now(),
    )
    session.add(row)
    session.sync.commit()
    return row


async def test_login_and_server_roles(secured, monkeypatch):
    import secrets

    password = secrets.token_urlsafe(24)
    client, _ = secured
    monkeypatch.setattr(
        get_settings(),
        "auth_users",
        {
            "operator": {
                "role": "coordinator",
                "password_hash": password_hash(password),
                "patient_ids": [],
            }
        },
    )
    correct = await client.post(
        "/api/v1/auth/login", json={"username": "operator", "password": password}
    )
    assert correct.status_code == 200
    token = correct.json()["access_token"]
    assert (
        await client.get(
            "/api/v1/coordinator/study-types", headers={"Authorization": "Bearer " + token}
        )
    ).status_code == 200
    assert (
        await client.post("/api/v1/auth/login", json={"username": "operator", "password": "wrong"})
    ).status_code == 401
    for path in ["/api/v1/coordinator/patients", "/api/v1/admin/triggers", "/api/v1/routes"]:
        assert (await client.get(path)).status_code == 401
    assert (
        await client.get("/api/v1/coordinator/patients", headers=auth_headers("doctor"))
    ).status_code == 403
    assert (
        await client.get("/api/v1/admin/triggers", headers=auth_headers("coordinator"))
    ).status_code == 403
    assert (
        await client.post(
            f"/api/v1/routes/{uuid4()}/tactics",
            json={"tactics": "observation"},
            headers=auth_headers("coordinator"),
        )
    ).status_code == 403


async def test_demo_login_without_setup(secured, monkeypatch):
    client, _ = secured
    monkeypatch.setattr(get_settings(), "environment", "dev")
    demo = (await client.get("/api/v1/auth/demo")).json()
    assert len(demo["roles"]) == 6
    for role in demo["roles"]:
        response = await client.post(
            "/api/v1/auth/login", json={"username": role, "password": demo["password"]}
        )
        assert response.status_code == 200
        assert response.json()["expires_in"] > 0
    monkeypatch.setattr(get_settings(), "environment", "production")
    assert (await client.get("/api/v1/auth/demo")).json()["enabled"] is False
    assert (
        await client.post(
            "/api/v1/auth/login", json={"username": "coordinator", "password": demo["password"]}
        )
    ).status_code == 401


async def test_tampered_and_expired_token(secured):
    import time

    from app.auth import sign

    client, _ = secured
    expired = sign({"sub": "test", "role": "admin", "exp": time.time() - 1})
    good = auth_headers()["Authorization"][7:]
    for token in [expired, good + "0", "bad-data", "a.b"]:
        assert (
            await client.get(
                "/api/v1/coordinator/study-types", headers={"Authorization": "Bearer " + token}
            )
        ).status_code == 401


@pytest.mark.parametrize("extension", ["txt", "docx"])
async def test_upload_preview_confirm_quotes_and_no_appointments(secured, extension):
    client, session = secured
    client.headers.update(auth_headers("coordinator"))
    text = (
        "ФИО: Иванов Иван Иванович\nТелефон: +7 (999) 123-45-67\n"
        "Адрес: ул. Лесная, дом 7\nДата рождения: 12.05.1980\n"
        "Заключение: Полип эндометрия 12 мм."
    )
    if extension == "docx":
        document = docx.Document()
        # Персональные данные в таблице тоже должны исчезнуть.
        document.add_paragraph("Заключение:" + text.split("Заключение:")[1])
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = text.split("Заключение:")[0]
        stream = io.BytesIO()
        document.save(stream)
        raw = stream.getvalue()
    else:
        raw = text.encode()
    response = await client.post(
        "/api/v1/coordinator/upload/preview",
        files={"file": ("protocol." + extension, raw)},
        data={"study_type": "УЗИ органов малого таза"},
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert "Иванов" not in preview["text"]
    assert "+7" not in preview["text"]
    assert "Лесная" not in preview["text"]
    assert "1980" not in preview["text"]
    assert len(preview["changes"]) >= 4
    assert list(session.sync.scalars(select(Study))) == []
    unreviewed = await client.post(
        "/api/v1/coordinator/upload/confirm", json={"ticket": preview["ticket"], "reviewed": False}
    )
    assert unreviewed.status_code == 403
    response = await client.post(
        "/api/v1/coordinator/upload/confirm", json={"ticket": preview["ticket"], "reviewed": True}
    )
    assert response.status_code == 201, response.text
    patient_id = response.json()["patient_id"]
    study = session.sync.scalar(select(Study))
    protocol = session.sync.scalar(select(Protocol))
    assert study and protocol
    assert "Иванов" not in study.raw_text
    assert "Иванов" not in str(protocol.facts)
    findings = session.sync.scalars(select(Finding)).all()
    assert findings
    for finding in findings:
        assert study.raw_text[finding.char_start : finding.char_end] == finding.quote
        assert finding.quote in text
    assert session.sync.scalars(select(Appointment)).all() == []
    card = await client.get(f"/api/v1/coordinator/patients/{patient_id}/protocols")
    assert card.status_code == 200
    assert card.json()["total"] == 1
    repeated = await client.post(
        "/api/v1/coordinator/upload/confirm", json={"ticket": preview["ticket"], "reviewed": True}
    )
    assert repeated.status_code == 201
    assert repeated.json()["study_id"] == response.json()["study_id"]
    assert len(session.sync.scalars(select(Study)).all()) == 1
    assert card.json()["items"][0]["rules"][0]["rule"]
    assert (await client.get("/api/v1/coordinator/patients")).json()["total"] == 1


async def test_search_all_protocols_and_foreign_patient(secured):
    client, session = secured
    demo = MockMisService().catalog()[0]
    patient = make_patient(session, demo["patient_id"], demo["external_id"])
    foreign = make_patient(session, card="foreign-card")
    for index in range(15):
        study = Study(
            id=uuid4(),
            patient_id=patient.id,
            study_type="УЗИ органов малого таза",
            study_date=get_clock().now().date() - timedelta(days=index),
            raw_text="Полип эндометрия",
            created_at=get_clock().now(),
        )
        session.add(study)
        session.sync.flush()
        session.add(
            Protocol(
                id=uuid4(),
                study_id=study.id,
                signed_at=get_clock().now(),
                created_at=get_clock().now(),
            )
        )
    session.sync.commit()
    client.headers.update(auth_headers("coordinator", [str(patient.id)]))
    for query in [demo["name"], demo["external_id"], str(study.id), demo["studies"][0]["study_id"]]:
        result = (await client.get("/api/v1/coordinator/patients", params={"q": query})).json()
        assert result["total"] == 1
    empty = (await client.get("/api/v1/coordinator/patients?q=неизвестный")).json()
    assert empty["total"] == 0
    assert "Измените запрос" in empty["message"]
    result = (await client.get(f"/api/v1/coordinator/patients/{patient.id}/protocols")).json()
    assert result["total"] == 15
    dates = [p["date"] for p in result["items"]]
    assert dates == sorted(dates, reverse=True)
    assert (
        await client.get(f"/api/v1/coordinator/patients/{foreign.id}/protocols")
    ).status_code == 403
    assert (await client.get(f"/api/v1/mock/lk/{foreign.id}/messages")).status_code == 403
    assert (
        await client.get(f"/api/v1/coordinator/patients/{uuid4()}/protocols")
    ).status_code == 404
    catalog = (await client.get("/api/v1/mock/mis/patients")).json()
    assert len(catalog) == 1
    assert catalog[0]["patient_id"] == str(patient.id)


async def test_confirm_foreign_patient_and_wrong_operator(secured):
    client, session = secured
    foreign = make_patient(session)
    client.headers.update(auth_headers("coordinator"))
    preview = (
        await client.post(
            "/api/v1/coordinator/upload/preview",
            files={"file": ("a.txt", b"empty")},
            data={"study_type": "УЗИ органов малого таза"},
        )
    ).json()
    assert (
        await client.post(
            "/api/v1/coordinator/upload/confirm",
            json={"ticket": preview["ticket"], "reviewed": True, "patient_id": str(foreign.id)},
        )
    ).status_code == 403
    # Билет предпросмотра не является токеном входа.
    assert (
        await client.get(
            "/api/v1/coordinator/patients", headers={"Authorization": "Bearer " + preview["ticket"]}
        )
    ).status_code == 401
    import time

    from app.auth import sign

    other = sign(
        {"sub": "other", "role": "coordinator", "patient_ids": [], "exp": time.time() + 100}
    )
    assert (
        await client.post(
            "/api/v1/coordinator/upload/confirm",
            json={"ticket": preview["ticket"], "reviewed": True},
            headers={"Authorization": "Bearer " + other},
        )
    ).status_code == 403


def test_anonymization_preserves_absolute_offsets():
    text = "Иванов Иван Иванович, +7 999 123 45 67, 1980-05-12\nЗаключение: Полип эндометрия 12 мм."
    clean, changes = anonymize(text)
    assert len(clean) == len(text)
    quote = "Полип эндометрия 12 мм."
    assert clean.index(quote) == text.index(quote)
    assert changes
    assert "Иванов" not in clean


async def test_lists_and_events_do_not_expand_scope(secured):
    from app.models import Route

    client, session = secured
    own = make_patient(session, card="own")
    foreign = make_patient(session, card="foreign")
    for patient in [own, foreign]:
        session.add(
            Route(id=uuid4(), patient_id=patient.id, status="created", created_at=get_clock().now())
        )
    session.sync.commit()
    routes = session.sync.scalars(select(Route)).all()
    client.headers.update(auth_headers("coordinator", [str(own.id)]))
    listing = (await client.get("/api/v1/routes")).json()
    assert listing["total"] == 1
    assert listing["items"][0]["patient_id"] == str(own.id)
    foreign_route = next(r for r in routes if r.patient_id == foreign.id)
    assert (await client.get("/api/v1/routes/" + str(foreign_route.id))).status_code == 403
    assert (
        await client.post(
            "/api/v1/mis/events",
            json={
                "event_id": "forbidden",
                "event_type": "StudyProtocolSigned",
                "occurred_at": get_clock().now().isoformat(),
                "subject": {"patient_id": str(foreign.id), "study_id": str(uuid4())},
                "payload": {"text": "test"},
            },
        )
    ).status_code == 403
    assert (
        await client.post(
            "/api/v1/mis/events",
            json={
                "event_id": "doctor-only",
                "event_type": "TacticsChosen",
                "occurred_at": get_clock().now().isoformat(),
                "subject": {"patient_id": str(own.id)},
                "payload": {"tactics": "observation"},
            },
        )
    ).status_code == 403


async def test_registered_screen_and_api(secured):
    client, _ = secured
    assert (await client.get("/coordinator")).status_code == 200
    for filename in ["auth.js", "coordinator.js", "coordinator.css"]:
        assert (await client.get("/static/" + filename)).status_code == 200
    paths = create_app().openapi()["paths"]
    for path in [
        "/patients",
        "/patients/{patient_id}/protocols",
        "/tasks",
        "/study-types",
        "/upload/preview",
        "/upload/confirm",
    ]:
        assert "/api/v1/coordinator" + path in paths


async def test_sort_by_name(secured):
    client, session = secured
    catalog = MockMisService().catalog()[:3]
    for patient in catalog:
        make_patient(session, patient["patient_id"], patient["external_id"])
    client.headers.update(auth_headers("coordinator", [str(p["patient_id"]) for p in catalog]))
    rows = (await client.get("/api/v1/coordinator/patients?sort=name")).json()["items"]
    assert [p["name"].casefold() for p in rows] == sorted(p["name"].casefold() for p in rows)
    assert (await client.get("/api/v1/coordinator/patients?sort=unknown")).status_code == 422
