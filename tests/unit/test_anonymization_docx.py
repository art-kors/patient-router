"""Обезличивание таблиц Word сохраняет клинику и смещения."""

from io import BytesIO

from docx import Document

from app.services.anonymization import anonymize_document, read_docx


def protocol(personal: bool = True) -> bytes:
    document = Document()
    document.add_paragraph("Заключение: «Полип эндометрия 12 мм».")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = (
        "ФИО: Иванов Иван Иванович\nТелефон: +7 (999) 123-45-67\n"
        "Адрес: ул. Лесная, дом 7\nДата рождения: 12.05.1980"
        if personal
        else "Пол: женский"
    )
    table.cell(0, 1).text = (
        "Возраст на момент осмотра: 46 лет\nВрач: Петров Петр Петрович\nДата приёма: 05.10.2026"
    )
    document.add_paragraph("Размеры: 12 мм; BI-RADS: 3; ORADS: 2")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_docx_table_personal_data() -> None:
    raw = protocol()
    text, changes = anonymize_document(raw)
    assert len(changes) >= 4
    for value in ("Иванов", "+7", "Лесная", "1980"):
        assert value not in text
    assert "Возраст на момент осмотра: 46 лет" in text
    assert "Врач: Петров Петр Петрович" in text
    assert "Дата приёма: 05.10.2026" in text
    original = read_docx(raw)
    for change in changes:
        start, end = change["start"], change["end"]
        assert original[start:end] == change["original"]
        assert text[start:end] == " " * (end - start)
    assert "1980\n\nВозраст" in original


def test_docx_conclusion_and_offsets_preserved() -> None:
    raw = protocol()
    original = read_docx(raw)
    text, _ = anonymize_document(raw)
    quote = "«Полип эндометрия 12 мм»"
    start = original.index(quote)
    assert text[start : start + len(quote)] == quote
    assert "Размеры: 12 мм; BI-RADS: 3; ORADS: 2" in text


def test_docx_length_preserved() -> None:
    raw = protocol()
    text, _ = anonymize_document(raw)
    assert len(text) == len(read_docx(raw))


def test_docx_without_personal_data_unchanged() -> None:
    raw = protocol(personal=False)
    text, changes = anonymize_document(raw)
    assert text == read_docx(raw)
    assert changes == []


def test_docx_label_and_value_in_separate_cells() -> None:
    """В настоящем протоколе дата и возраст стоят отдельно от меток."""
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Дата\xa0рождения:"
    table.cell(0, 1).text = "31.07.1987"
    table.cell(1, 0).text = "Возраст\xa0на\xa0момент\xa0осмотра:"
    table.cell(1, 1).text = "39\xa0лет"
    document.add_paragraph("Заключение: Полип эндометрия 12 мм")
    stream = BytesIO()
    document.save(stream)
    raw = stream.getvalue()
    text, changes = anonymize_document(raw)
    assert "31.07.1987" not in text
    assert "39\xa0лет" in text
    assert "Полип эндометрия 12 мм" in text
    assert len(text) == len(read_docx(raw))
    assert len(changes) == 1
